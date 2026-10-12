import io
from docx import Document
from .pdf_parser import clean_extracted_text, detect_resume_metadata

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract full text from DOCX bytes including paragraphs and table contents."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        
        # Read normal paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)
                
        # Read table cells as well (many resumes use tables for layouts)
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    # Deduplicate adjoining repeated cells from merged rows
                    unique_texts = []
                    for t in row_texts:
                        if not unique_texts or unique_texts[-1] != t:
                            unique_texts.append(t)
                    paragraphs.append(" | ".join(unique_texts))

        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("The uploaded DOCX document contains no readable text.")
            
        return clean_extracted_text(full_text)
    except Exception as e:
        raise ValueError(f"Failed to read DOCX file: {str(e)}")
