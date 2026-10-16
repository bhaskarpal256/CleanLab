import io
from typing import Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_horizontal_rule(paragraph):
    """Add a subtle horizontal border below a section header."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '4F46E5')  # Indigo accent
    pBdr.append(bottom)
    pPr.append(pBdr)

def generate_docx(resume_data: Dict[str, Any]) -> bytes:
    """Generate a clean, executive ATS-friendly DOCX document optimized to fit on a single page."""
    doc = Document()
    
    # Page Margins (0.45 in for executive 1-page fit)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(31, 41, 55) # Dark gray #1F2937

    # 1. Candidate Name (Header)
    name_paragraph = doc.add_paragraph()
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_paragraph.add_run(resume_data.get("candidate_name", "Candidate Name"))
    name_run.font.size = Pt(18)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(17, 24, 39)
    name_paragraph.paragraph_format.space_before = Pt(0)
    name_paragraph.paragraph_format.space_after = Pt(1)

    # 2. Contact Information
    contact = resume_data.get("contact_info", {})
    contact_parts = [
        contact.get("email"),
        contact.get("phone"),
        contact.get("location"),
        contact.get("linkedin"),
        contact.get("github"),
        contact.get("portfolio")
    ]
    valid_contact_parts = [str(p) for p in contact_parts if p]
    if valid_contact_parts:
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_run = contact_para.add_run("  •  ".join(valid_contact_parts))
        contact_run.font.size = Pt(9.0)
        contact_run.font.color.rgb = RGBColor(75, 85, 99)
        contact_para.paragraph_format.space_after = Pt(6)

    def add_section_title(title_text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title_text.upper())
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
        add_horizontal_rule(p)
        return p

    # 3. Summary
    summary_text = resume_data.get("summary")
    if summary_text:
        add_section_title("Professional Summary")
        summary_para = doc.add_paragraph(summary_text)
        summary_para.paragraph_format.space_after = Pt(3)
        summary_para.paragraph_format.line_spacing = 1.05

    # 4. Technical Skills
    skill_groups = resume_data.get("skill_groups", [])
    if skill_groups:
        add_section_title("Technical Skills & Core Competencies")
        for group in skill_groups:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            run_cat = p.add_run(f"{group.get('category')}: ")
            run_cat.bold = True
            run_cat.font.size = Pt(9.5)
            run_cat.font.color.rgb = RGBColor(31, 41, 55)
            run_skills = p.add_run(", ".join(group.get("skills", [])))
            run_skills.font.size = Pt(9.5)

    # 5. Professional Experience
    experiences = resume_data.get("experience", [])
    if experiences:
        add_section_title("Professional Experience")
        for exp in experiences:
            p_job = doc.add_paragraph()
            p_job.paragraph_format.space_before = Pt(4)
            p_job.paragraph_format.space_after = Pt(1)
            
            run_role = p_job.add_run(exp.get("role", "Software Engineer"))
            run_role.bold = True
            run_role.font.size = Pt(10)
            
            run_comp = p_job.add_run(f" | {exp.get('company', 'Company')}")
            run_comp.font.size = Pt(10)
            
            details = []
            if exp.get("location"):
                details.append(exp.get("location"))
            if exp.get("duration"):
                details.append(exp.get("duration"))
            if details:
                run_meta = p_job.add_run(f" ({' - '.join(details)})")
                run_meta.italic = True
                run_meta.font.size = Pt(9.0)
                run_meta.font.color.rgb = RGBColor(107, 114, 128)

            bullets = exp.get("tailored_bullets", []) or exp.get("original_bullets", [])
            for bullet in bullets:
                p_bullet = doc.add_paragraph(style='List Bullet')
                p_bullet.paragraph_format.space_before = Pt(0)
                p_bullet.paragraph_format.space_after = Pt(1.5)
                p_bullet.paragraph_format.line_spacing = 1.05
                run_b = p_bullet.add_run(bullet)
                run_b.font.size = Pt(9.5)

    # 6. Key Projects
    projects = resume_data.get("projects", [])
    if projects:
        add_section_title("Projects & Technical Implementations")
        for p_item in projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(4)
            p_proj.paragraph_format.space_after = Pt(1)
            
            run_title = p_proj.add_run(p_item.get("title", "Project"))
            run_title.bold = True
            run_title.font.size = Pt(10)
            
            techs = p_item.get("technologies", [])
            if techs:
                run_tech = p_proj.add_run(f" [{', '.join(techs)}]")
                run_tech.font.size = Pt(9.0)
                run_tech.font.color.rgb = RGBColor(107, 114, 128)
                
            desc = p_item.get("description")
            if desc:
                p_desc = doc.add_paragraph(desc)
                p_desc.paragraph_format.space_after = Pt(1)
                p_desc.paragraph_format.line_spacing = 1.05
                
            for b in p_item.get("tailored_bullets", []):
                p_b = doc.add_paragraph(style='List Bullet')
                p_b.paragraph_format.space_before = Pt(0)
                p_b.paragraph_format.space_after = Pt(1.5)
                p_b.paragraph_format.line_spacing = 1.05
                run_pb = p_b.add_run(b)
                run_pb.font.size = Pt(9.5)

    # 7. Education
    education = resume_data.get("education", [])
    if education:
        add_section_title("Education")
        for edu in education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(1)
            p_edu.paragraph_format.space_after = Pt(1.5)
            deg_run = p_edu.add_run(edu.get("degree", "Degree"))
            deg_run.bold = True
            deg_run.font.size = Pt(9.5)
            inst_run = p_edu.add_run(f" — {edu.get('institution', 'University')}")
            inst_run.font.size = Pt(9.5)
            if edu.get("year"):
                yr_run = p_edu.add_run(f" ({edu.get('year')})")
                yr_run.font.size = Pt(9.0)
                yr_run.font.color.rgb = RGBColor(107, 114, 128)
            if edu.get("gpa_or_honors"):
                gpa_run = p_edu.add_run(f" • {edu.get('gpa_or_honors')}")
                gpa_run.font.size = Pt(9.0)
                gpa_run.font.color.rgb = RGBColor(107, 114, 128)

    # 8. Academic Achievements & Leadership Experience
    achievements = resume_data.get("achievements", [])
    if achievements:
        add_section_title("Academic Achievements & Leadership")
        for ach in achievements:
            p_ach = doc.add_paragraph(style='List Bullet')
            p_ach.paragraph_format.space_before = Pt(0)
            p_ach.paragraph_format.space_after = Pt(1.5)
            p_ach.paragraph_format.line_spacing = 1.05
            run_ach = p_ach.add_run(ach)
            run_ach.font.size = Pt(9.0)

    output_stream = io.BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    return output_stream.getvalue()
